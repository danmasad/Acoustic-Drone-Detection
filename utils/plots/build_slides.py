"""Build the talk (docs/presentation.pptx) and the two figures in the report.

Numbers are the ones in docs/report.md (12 October 2026). The headline bars are
checked against runs/grid/*/metrics.json before anything is written.

    MPLCONFIGDIR=/tmp/mpl .venv/bin/python utils/plots/build_slides.py
"""
import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
FIG = ROOT / "figures"
DOCS = ROOT / "docs"
OUT = DOCS / "presentation.pptx"

FONT = "Helvetica Neue"
BG, DARK = "FAFAF9", "0A0A0A"
INK, BODY, MUTED, LINE, CARD = "1C1917", "44403C", "78716C", "E7E5E4", "FFFFFF"
BLUE, GREEN, RED, AMBER, GREY = "2563EB", "16A34A", "DC2626", "D97706", "A8A29E"
SW, SH = 13.333, 7.5
TOTAL = 14

# Seed means on the matched site, before the filter. Delay-and-sum is one run;
# 32.6% rounds to the 33 in the report table.
HEADLINE = [
    ("Delay-and-sum, 24", 32.6, 0.0, GREY),
    ("24 mics, 1 frame", 70.5, 1.5, BLUE),
    ("Device 3, 1 frame", 69.3, 0.7, GREEN),
    ("24 mics, 9 frames", 70.1, 2.7, BLUE),
    ("Device 3, 9 frames", 73.2, 1.4, GREEN),
    ("Searched 6, 1 frame", 52.1, 0.8, AMBER),
]
# Same device-3 K=9 weights; only the depth of real history changes.
HISTORY_GAIN = [0.5, 4.4, 8.7, 21.8, 3.0]
HISTORY_BINS = ["0–30 m", "30–60", "60–90", "90–120", "120–180"]


def rgb(h):
    return RGBColor.from_string(h)


def box(slide, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    if radius is not None:
        s.adjustments[0] = radius
    return s


def card(slide, x, y, w, h, fill=CARD, line=LINE):
    return box(slide, x, y, w, h, fill, line, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)


def text(slide, x, y, w, h, paras, size=16, color=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, after=6, line=1.08):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        p.line_spacing = line
        runs = [(para, {})] if isinstance(para, str) else para
        for t, st in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", bold)
            f.italic = st.get("italic", False)
            f.color.rgb = rgb(st.get("color", color))
    return tb


def bullets(slide, x, y, w, h, items, size=16, dot=BLUE, after=8):
    paras = []
    for it in items:
        runs = [("●  ", {"color": dot, "size": size - 5})]
        runs += [(it, {})] if isinstance(it, str) else it
        paras.append(runs)
    return text(slide, x, y, w, h, paras, size=size, after=after, line=1.08)


def hline(slide, x1, y1, x2, y2, color=LINE, width=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    return c


def base(prs, n, kicker, title):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    if kicker:
        text(s, 0.55, 0.32, 12, 0.28, kicker.upper(), size=12, color=BLUE, bold=True)
    if title:
        text(s, 0.55, 0.58, 12.2, 0.7, title, size=28, color=INK, bold=True, after=0)
    if n:
        text(s, 0.55, 7.12, 9.2, 0.24, "Acoustic Drone Detection  ·  Shai Livshits & Dan Masad", size=11, color=MUTED)
        text(s, 11.3, 7.12, 1.5, 0.24, f"{n}  /  {TOTAL}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    return s


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def bottom_line(slide, y, body, h=0.9):
    card(slide, 0.55, y, 12.2, h, fill=DARK, line=None)
    text(slide, 0.8, y + 0.16, 2.0, 0.3, "BOTTOM LINE", size=12, color="60A5FA", bold=True)
    text(slide, 2.65, y + 0.12, 9.9, h - 0.2, body, size=16, color="F5F5F4", line=1.08, after=0)


def stat(slide, x, y, w, h, big, label, color=INK, big_size=26):
    card(slide, x, y, w, h)
    text(slide, x + 0.18, y + 0.14, w - 0.36, 0.5, big, size=big_size, color=color, bold=True)
    text(slide, x + 0.18, y + 0.68, w - 0.36, h - 0.82, label, size=13, color=BODY, after=0)


def table(slide, x, y, widths, headers, rows, size=14, rh=0.42):
    cx = x
    for j, h in enumerate(headers):
        text(slide, cx, y, widths[j] - 0.08, 0.36, h, size=12, color=MUTED, bold=True,
             align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER, after=0)
        cx += widths[j]
    hline(slide, x, y + 0.40, x + sum(widths), y + 0.40, color=MUTED, width=1.0)
    for i, row in enumerate(rows):
        yy = y + 0.50 + i * rh
        cx = x
        for j, cell in enumerate(row):
            text(slide, cx, yy, widths[j] - 0.08, 0.34, cell, size=size, color=INK,
                 bold=(j == 0), align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
            cx += widths[j]
        if i < len(rows) - 1:
            hline(slide, x, yy + rh - 0.06, x + sum(widths), yy + rh - 0.06, width=0.6)
    return y + 0.50 + len(rows) * rh


def _style(ax):
    ax.set_facecolor("#FAFAF9")
    for sp in ax.spines.values():
        sp.set_color("#E7E5E4")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(colors="#44403C", labelsize=12, length=0)
    ax.grid(axis="x", color="#E7E5E4", zorder=0)


def check_headline() -> None:
    """The bars have to be the runs, not a retyped copy of the report."""
    expected = {("all24", 1): (70.5, 1.5), ("all24", 9): (70.1, 2.7),
                ("device3", 1): (69.3, 0.7), ("device3", 9): (73.2, 1.4),
                ("best6", 1): (52.1, 0.8)}
    for (subset, k), (mean, sd) in expected.items():
        vals = []
        for seed in range(3):
            path = ROOT / "runs" / "grid" / f"{subset}_K{k}_w32_b16_wce_s{seed}" / "metrics.json"
            vals.append(json.loads(path.read_text())["test_2"]["hit_10deg"] * 100)
        got_m, got_s = float(np.mean(vals)), float(np.std(vals, ddof=1))
        if abs(got_m - mean) > 0.15 or abs(got_s - sd) > 0.15:
            raise SystemExit(f"{subset} K={k}: runs say {got_m:.2f} ± {got_s:.2f}, "
                             f"report says {mean} ± {sd}")


def save_charts() -> dict:
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Helvetica Neue", "Helvetica", "Arial"],
        "axes.unicode_minus": False,
    })
    check_headline()

    fig, ax = plt.subplots(figsize=(7.4, 3.55), dpi=160)
    fig.patch.set_facecolor("#FAFAF9")
    labels = [r[0] for r in HEADLINE][::-1]
    means = [r[1] for r in HEADLINE][::-1]
    sds = [r[2] for r in HEADLINE][::-1]
    colors = ["#" + r[3] for r in HEADLINE][::-1]
    y = np.arange(len(labels))
    ax.barh(y, means, xerr=sds, color=colors, height=0.62, zorder=2,
            error_kw={"ecolor": "#44403C", "capsize": 3, "capthick": 1, "elinewidth": 1})
    for i, (m, s) in enumerate(zip(means, sds)):
        label = f"{m:.1f}" if s == 0 else f"{m:.1f} ± {s:.1f}"
        ax.text(m + s + 1.3, i, label, va="center", ha="left", color="#1C1917", fontsize=12)
    ax.set_yticks(y)
    ax.set_yticklabels(labels)
    ax.set_xlim(0, 100)
    ax.set_xlabel("Share of drone frames within 10°")
    _style(ax)
    fig.tight_layout()
    headline = DOCS / "headline_within_10.png"
    fig.savefig(headline, facecolor=fig.get_facecolor())
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7.4, 3.15), dpi=160)
    fig.patch.set_facecolor("#FAFAF9")
    xs = np.arange(len(HISTORY_GAIN))
    colors = ["#16A34A" if g >= 8 else "#2563EB" if g >= 2 else "#A8A29E" for g in HISTORY_GAIN]
    ax.bar(xs, HISTORY_GAIN, color=colors, width=0.66, zorder=2)
    for i, g in enumerate(HISTORY_GAIN):
        ax.text(i, g + 0.6, f"+{g:.1f}", ha="center", va="bottom", color="#1C1917", fontsize=13)
    ax.set_xticks(xs)
    ax.set_xticklabels(HISTORY_BINS)
    ax.set_ylim(0, 28)
    ax.set_ylabel("Gain, points within 10°")
    _style(ax)
    ax.grid(axis="y", color="#E7E5E4", zorder=0)
    ax.grid(axis="x", visible=False)
    fig.tight_layout()
    history = DOCS / "history_gain.png"
    fig.savefig(history, facecolor=fig.get_facecolor())
    plt.close(fig)

    runs = []
    for seed in range(3):
        path = ROOT / "runs" / "grid" / f"device3_K1_w32_b16_wce_s{seed}" / "metrics.json"
        m = json.loads(path.read_text())
        runs.append((m["best_epoch"], m["history"]))
    fig, (ax0, ax1) = plt.subplots(1, 2, figsize=(12.0, 3.35), dpi=150)
    fig.patch.set_facecolor("#FAFAF9")
    palette = ["#2563EB", "#0F766E", "#D97706"]
    for seed, ((best, hist), color) in enumerate(zip(runs, palette)):
        ep = [h["epoch"] for h in hist]
        ax0.plot(ep, [h["hit_10deg"] * 100 for h in hist], color=color, lw=1.8, label=f"seed {seed}")
        ax0.scatter([best], [hist[best]["hit_10deg"] * 100], s=36, color=color, zorder=3)
        ax1.plot(ep, [h["train_loss"] for h in hist], color=color, lw=1.6)
        ax1.plot(ep, [h["val_loss"] for h in hist], color=color, lw=1.4, ls="--")
    ax0.set_ylim(52, 78)
    ax0.set_ylabel("Within 10° on test 2 (%)")
    ax0.set_title("The direction. This holds.", loc="left", fontsize=13, color="#1C1917")
    ax0.legend(frameon=False, fontsize=11)
    ax1.set_ylabel("Pixel loss")
    ax1.set_title("The loss. This is not the score.", loc="left", fontsize=13, color="#1C1917")
    for ax in (ax0, ax1):
        _style(ax)
        ax.set_xlabel("Epoch")
        ax.grid(axis="y", color="#E7E5E4", zorder=0)
        ax.grid(axis="x", visible=False)
    fig.tight_layout(w_pad=3.0)
    training = DOCS / "_training_device3_k1.png"
    fig.savefig(training, facecolor=fig.get_facecolor())
    plt.close(fig)
    return {"headline": headline, "history": history, "training": training}


def s01_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(DARK)
    text(s, 0.7, 1.55, 6.2, 0.3, "DEEP LEARNING 2026B  ·  APPLIED", size=13, color="60A5FA", bold=True)
    text(s, 0.7, 2.05, 6.2, 1.8, ["Acoustic Drone", "Detection"], size=44, color="F5F5F4", bold=True, after=0, line=0.96)
    text(s, 0.7, 4.35, 5.9, 0.32, "Shai Livshits  ·  Dan Masad", size=16, color="F5F5F4")
    text(s, 0.7, 4.78, 5.9, 0.28, "12 October 2026", size=14, color="78716C")
    picture(s, FIG / "flight_hemisphere_mar18_f2.gif", 7.15, 1.7, w=5.4)


def s02_problem(prs):
    s = base(prs, 2, "The problem", "A drone on a wire is silent to every sensor but one")
    text(s, 0.55, 1.42, 12.2, 0.3,
         "Fibre-optic FPV drones radiate nothing.", size=15, color=BODY)
    hurdles = [
        ("Radio", "Nothing to hear", "No link. Nothing to scan, nothing to jam."),
        ("Radar", "Too small, too low", "Plastic, sub-metre, in the ground clutter."),
        ("Camera", "Too late", "A few pixels. Gone at night or in fog."),
    ]
    x = 0.55
    for name, line1, body in hurdles:
        card(s, x, 1.95, 3.85, 1.9)
        text(s, x + 0.22, 2.1, 3.4, 0.3, name.upper(), size=12, color=RED, bold=True)
        text(s, x + 0.22, 2.42, 3.4, 0.5, line1, size=20, color=INK, bold=True, after=0)
        text(s, x + 0.22, 3.05, 3.4, 0.7, body, size=14, color=BODY, line=1.12)
        x += 4.0
    card(s, 0.55, 4.2, 12.2, 2.65, fill=DARK, line=None)
    text(s, 0.8, 4.4, 5.4, 0.3, "WHAT IT CANNOT SWITCH OFF", size=12, color="60A5FA", bold=True)
    text(s, 0.8, 4.78, 5.4, 1.9,
         ["Rotors are loud, tonal, and heard from every direction.",
          "Listening is passive."],
         size=18, color="F5F5F4", line=1.15, after=10)
    text(s, 6.6, 4.4, 5.9, 0.3, "OUR METHOD", size=12, color="60A5FA", bold=True)
    steps = [
        "Microphones hear the same sound with small delays.",
        "Beamforming turns the delays into a map of the sky.",
        "A network marks the drone on the map, 10× a second.",
    ]
    tb = bullets(s, 6.6, 4.78, 5.9, 1.9, steps, size=16, dot="60A5FA", after=8)
    for p in tb.text_frame.paragraphs:
        for r in p.runs[1:]:                 # keep the blue dot, lighten the body text
            r.font.color.rgb = rgb("F5F5F4")


def s03_data(prs):
    s = base(prs, 3, "Data", "One drone, two sites, and borrowed quiet")
    text(s, 0.55, 1.42, 12.2, 0.3,
         "Western Sydney University dataset. DJI Air 3, 24 microphones, 48 kHz, GPS at 10 Hz.",
         size=15, color=BODY)
    table(s, 0.55, 1.9, [2.0, 2.4, 2.6, 5.2],
          ["Split", "Drone audio", "Where", "Role"],
          [
              ["Train", "52 min", "Site 1", "Fit the weights"],
              ["Test 2", "20 min", "Site 1", "Pick the epoch"],
              ["Test 1", "3 min", "Second site", "Sealed. Scored once"],
          ], size=16, rh=0.5)
    text(s, 0.55, 4.0, 8, 0.26, "NO-DRONE AUDIO: 5 MINUTES IN THE RELEASE", size=12, color=BLUE, bold=True)
    quiet = [
        ("Pre-takeoff", "4.8 → 8.2 min", "Mined from the same recordings, before the motors spin up."),
        ("ESC-50", "50 classes", "Short everyday sounds, mixed into our quiet frames."),
        ("DEMAND", "18 places", "90 min of rooms, streets, vehicles. Train on five kinds, score the sixth."),
    ]
    x = 0.55
    for name, big, body in quiet:
        card(s, x, 4.38, 3.85, 2.47)
        text(s, x + 0.22, 4.52, 3.4, 0.3, name.upper(), size=12, color=MUTED, bold=True)
        text(s, x + 0.22, 4.84, 3.4, 0.5, big, size=22, color=INK, bold=True, after=0)
        text(s, x + 0.22, 5.5, 3.4, 1.2, body, size=14, color=BODY, line=1.12)
        x += 4.0


def s06_model(prs):
    s = base(prs, 6, "Method", "What we tried")
    text(s, 0.55, 1.38, 6, 0.26, "WHERE WE STARTED", size=12, color=BLUE, bold=True)
    basis = [
        ("The map", "Delay-and-sum. Each 100 ms frame becomes 23 by 91 directions, in 16 bands."),
        ("The U-Net", "The paper's network. It reads that map as an image and paints a mask. The direction is read off the mask."),
        ("The baseline", "The brightest cell of the map itself. No learning, so a gain over it is the network."),
    ]
    for i, (title, body) in enumerate(basis):
        x = 0.55 + i * 4.2
        card(s, x, 1.66, 4.02, 1.5)
        text(s, x + 0.16, 1.78, 3.7, 0.32, title, size=16, color=BLUE, bold=True)
        text(s, x + 0.16, 2.14, 3.7, 0.95, body, size=14, color=BODY, line=1.06)
    text(s, 0.55, 3.34, 8, 0.26, "HOW WE TRIED TO GET AN EDGE", size=12, color=INK, bold=True)
    model_row = [
        ("History", "K past frames stacked as channels. K is 1, 5, or 9."),
        ("Six-mic subsets", "A greedy sidelobe search, plus simple picks. Device 3 matches all 24."),
        ("Two losses", "The paper's Tversky against weighted cross-entropy. 63.5% vs 70.1% within 10°."),
        ("Disc or grid", "A polar disc, zenith in the centre. 69.5% vs 70.1%, so we kept the grid."),
    ]
    for i, (title, body) in enumerate(model_row):
        x = 0.55 + i * 3.2
        card(s, x, 3.64, 3.05, 1.5)
        text(s, x + 0.14, 3.76, 2.76, 0.32, title, size=15, color=INK, bold=True)
        text(s, x + 0.14, 4.12, 2.76, 0.95, body, size=13, color=BODY, line=1.05)
    other_row = [
        ("Detection, two ways", "A presence score on the U-Net. And an 86k network on one microphone, with no array."),
        ("More quiet audio", "Pre-takeoff, 4.8 to 8.2 minutes. ESC-50 sounds mixed in. DEMAND places: train on five, score the sixth."),
        ("Kinematic filter", "After the U-Net. Based on drone speed limits. Not trained."),
    ]
    for i, (title, body) in enumerate(other_row):
        x = 0.55 + i * 4.2
        card(s, x, 5.3, 4.02, 1.5)
        text(s, x + 0.16, 5.42, 3.7, 0.32, title, size=15, color=INK, bold=True)
        text(s, x + 0.16, 5.78, 3.7, 0.95, body, size=13, color=BODY, line=1.05)


def save_unet(path: Path) -> None:
    """The trained U-Net: depth 3, base 32, mask head plus the presence head."""
    from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

    fig, ax = plt.subplots(figsize=(13.0, 5.35), dpi=150)
    fig.patch.set_facecolor("#FAFAF9")
    ax.set_xlim(0, 13)
    ax.set_ylim(0, 5.35)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_facecolor("#FAFAF9")

    def box(x, y, w, h, fc, ec, title, sub, tc="#1C1917", sc="#44403C", ts=15, ss=11):
        ax.add_patch(FancyBboxPatch(
            (x - w / 2, y - h / 2), w, h,
            boxstyle="round,pad=0.012,rounding_size=0.08",
            facecolor=fc, edgecolor=ec, linewidth=1.6, mutation_aspect=0.55, zorder=3))
        ax.text(x, y + (0.13 if sub else 0), title, ha="center", va="center",
                fontsize=ts, color=tc, fontweight="medium", zorder=4)
        if sub:
            ax.text(x, y - 0.2, sub, ha="center", va="center", fontsize=ss, color=sc, zorder=4)

    def arrow(x1, y1, x2, y2, color="#2563EB", lw=1.5, style="-"):
        ax.add_patch(FancyArrowPatch(
            (x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=12,
            color=color, linewidth=lw, linestyle=style, zorder=2,
            shrinkA=0, shrinkB=0))

    # Encoder left, decoder right. Sizes are what max-pooling a 23×91 map actually yields.
    xl, xr, w, h = 2.55, 7.55, 2.15, 0.78
    ys = [4.35, 3.15, 1.95]
    enc = [("#DBEAFE", "#2563EB", "32", "23 × 91"),
           ("#DBEAFE", "#2563EB", "64", "11 × 45"),
           ("#DBEAFE", "#2563EB", "128", "5 × 22")]
    dec = [("#DCFCE7", "#16A34A", "32", "23 × 91"),
           ("#DCFCE7", "#16A34A", "64", "11 × 45"),
           ("#DCFCE7", "#16A34A", "128", "5 × 22")]
    for y, (fc, ec, title, sub) in zip(ys, enc):
        box(xl, y, w, h, fc, ec, title, sub)
    for y, (fc, ec, title, sub) in zip(ys, dec):
        box(xr, y, w, h, fc, ec, title, sub)

    box(5.05, 0.72, 2.15, 0.78, "#1D4ED8", "#1D4ED8", "256", "2 × 11", tc="#F8FAFC", sc="#DBEAFE", ss=11)

    box(0.95, 4.35, 1.55, 0.7, "#FFFFFF", "#D6D3D1", "16 × K", "23 × 91", ts=13, ss=10)
    box(9.55, 4.35, 1.7, 0.7, "#FFFFFF", "#16A34A", "Mask", "23 × 91", ts=13, ss=10)
    box(11.85, 4.35, 1.85, 0.7, "#FFFFFF", "#44403C", "Direction", "brightest cell", ts=13, ss=10)

    box(9.15, 0.72, 2.05, 0.78, "#FEF3C7", "#D97706", "Avg + max", "512, then 128", ts=13, ss=10)
    box(11.7, 0.72, 1.9, 0.78, "#FEF3C7", "#D97706", "Presence", "drone in frame?", ts=13, ss=10)

    arrow(1.73, 4.35, xl - w / 2, 4.35)
    arrow(xr + w / 2, 4.35, 8.7, 4.35, color="#16A34A")
    arrow(10.4, 4.35, 10.92, 4.35, color="#44403C")
    for y0, y1 in zip(ys, ys[1:]):
        arrow(xl, y0 - h / 2, xl, y1 + h / 2)
        arrow(xr, y1 + h / 2, xr, y0 - h / 2, color="#16A34A")
    arrow(xl, ys[-1] - h / 2, 5.05, 0.72 + 0.39)
    arrow(5.05, 0.72 + 0.39, xr, ys[-1] - h / 2, color="#16A34A")
    for y in ys:
        arrow(xl + w / 2, y, xr - w / 2, y, color="#93C5FD", lw=1.3)
    arrow(6.13, 0.72, 8.12, 0.72, color="#D97706")
    arrow(10.18, 0.72, 10.75, 0.72, color="#D97706")

    ax.text(xl + 0.28, 3.78, "pool", fontsize=9, color="#2563EB", ha="left")
    ax.text(8.78, 3.75, "nearest + skip", fontsize=9, color="#16A34A", ha="left", va="center")
    ax.text(5.05, 2.55, "skip", fontsize=10, color="#60A5FA", ha="center")
    ax.text(6.5, 5.05, "Each box is two 3×3 convolutions, batch norm, and ReLU. Azimuth wraps, so the left and right edges are one bearing.",
            fontsize=11, color="#44403C", ha="center")
    fig.subplots_adjust(0, 0, 1, 1)
    fig.savefig(path, facecolor=fig.get_facecolor())
    plt.close(fig)


def s07_network(prs):
    s = base(prs, 7, "The network", "One U-Net, two answers")
    picture(s, DOCS / "unet_structure.png", 0.2, 1.35, w=12.9)


def s04_hardware(prs):
    s = base(prs, 4, "The array", "Four recorders, six microphones each")
    picture(s, FIG / "array_geometry.png", 0.28, 1.38, w=8.15)
    points = [
        ("Colour is the recorder", "Purple, blue, teal, yellow. Each is one Zoom recorder, so the six share a clock."),
        ("Teal is device 3", "The six we train. 1.09 m across, the outer microphones of one recorder."),
        ("The bottom row", "Other subsets we drew. Centre is device 3 again."),
    ]
    for i, (title, body) in enumerate(points):
        y = 1.48 + i * 1.3
        text(s, 8.6, y, 4.2, 0.32, title, size=16, color=INK, bold=True)
        text(s, 8.6, y + 0.34, 4.2, 0.85, body, size=14, color=BODY, line=1.05)


def s05_map(prs):
    s = base(prs, 5, "The map", "What the network sees")
    picture(s, FIG / "array_beampattern.png", 0.4, 1.4, w=9.2)
    points = [
        ("One source", "Every other bright spot is a phantom made by the array."),
        ("Why a network", "The brightest pixel can be a phantom. The network reads the whole pattern."),
    ]
    for i, (title, body) in enumerate(points):
        y = 2.2 + i * 1.9
        text(s, 9.9, y, 2.9, 0.34, title, size=18, color=INK, bold=True)
        text(s, 9.9, y + 0.42, 2.9, 1.3, body, size=15, color=BODY, line=1.08)


def s08_training(prs, charts):
    s = base(prs, 8, "Training", "We keep the best test-2 hit rate, not the last epoch")
    stats = [
        ("Adam, 5×10⁻³", "Constant. No weight decay. Batch 32. Batch norm, dropout 0.1."),
        ("Weighted loss", "1.4% of pixels are the drone. Unweighted, empty sky wins."),
        ("Patience 10", "At least 15 epochs, at most 60. Three seeds. Test 1 stays sealed."),
    ]
    for i, (big, label) in enumerate(stats):
        stat(s, 0.55 + i * 4.2, 1.45, 4.02, 1.35, big, label, color=BLUE, big_size=22)
    picture(s, charts["training"], 0.5, 3.0, w=12.3)
    text(s, 0.55, 6.48, 12.2, 0.55,
         "Device 3, one frame. The dot is the epoch we keep. Solid is training loss, dashed is test 2. The direction stays near 70%. The loss rises because a few confident wrong pixels dominate it.",
         size=14, color=BODY)


def s10_history(prs):
    s = base(prs, 10, "Results  ·  our benchmark", "History stands in for microphones")
    text(s, 0.55, 1.42, 12.2, 0.3,
         "Within 10°. Left: gain from one frame to nine. Right: device 3 at nine frames, all 24 at one.",
         size=15, color=BODY)
    picture(s, DOCS / "analysis_history.png", 1.0, 1.9, w=11.3)
    bottom_line(s, 5.95, "Nine frames lift device 3 to the level of all 24, mostly past 60 m. "
                "A sidelobe flickers from frame to frame; the true peak stays put.")


def s11_filter_gain(prs):
    s = base(prs, 11, "Results  ·  our benchmark", "The filter helps most without history")
    text(s, 0.55, 1.42, 12.2, 0.3,
         "Points gained within 10° when the kinematic filter is applied after the network. Test 2.",
         size=15, color=BODY)
    picture(s, DOCS / "analysis_filter.png", 1.3, 1.85, w=10.6)
    text(s, 0.55, 5.6, 12.2, 0.3, [
        [("All ranges:   ", {"bold": True, "color": INK}),
         ("24 mics, 1 frame  70 → 75%   ·   Device 3, 1 frame  69 → 73%   ·   Device 3, 9 frames  73 → 73%", {})],
    ], size=15, color=BODY, align=PP_ALIGN.CENTER)
    bottom_line(s, 6.05, "Highly effective at long range and when the network has no temporal context. "
                "Nine frames and the filter use the same window, so they do not add up.", h=0.85)


def s12_filter(prs):
    s = base(prs, 12, "After the network", "Inside 10°, we follow at once")
    picture(s, DOCS / "filter_sky.gif", 0.35, 1.42, w=6.4)
    card(s, 6.95, 1.42, 5.85, 3.4)
    text(s, 7.18, 1.6, 5.45, 0.34, "The rule", size=18, color=INK, bold=True)
    bullets(s, 7.18, 2.1, 5.45, 2.5, [
        "A drone has a top speed, so its direction can only change so much between consecutive frames.",
        "A new direction within 10° of the last one is taken at once. A larger jump is taken only after half a second of agreement.",
    ], size=17, after=10)


def s09_paper(prs):
    s = base(prs, 9, "Results  ·  the paper's benchmark", "Against the paper, on their metric")
    text(s, 0.55, 1.42, 12.2, 0.55,
         "Mean angular error, matched site, the paper's distance bins. U-Net rows are stable frames; "
         "delay-and-sum rows are every frame. Our U-Net cells are the mean of three seeds.",
         size=15, color=BODY)
    table(s, 0.55, 2.1, [4.6, 1.6, 2.0, 2.0, 2.0],
          ["", "Site", "0–50 m", "50–100 m", "100–200 m"],
          [
              ["Paper, delay-and-sum", "matched", "38°", "75°", "33°"],
              ["Ours, delay-and-sum, 24", "matched", "40°", "72°", "31°"],
              ["Paper, U-Net, stable", "matched", "10°", "13°", "14°"],
              ["Ours, U-Net 24, stable", "matched", "9°", "11°", "46°"],
              ["Ours, device 3, stable", "matched", "9°", "11°", "49°"],
          ], size=16, rh=0.5)
    bottom_line(s, 5.2, "We reproduce their performance under 100 m, where the flight is; they lead past 100 m. "
                "On this metric device 3 matches all 24.")
    text(s, 0.55, 6.3, 12.2, 0.3,
         "This slide is the paper's benchmark. The next slides score every frame on our own benchmark: the share within 10°.",
         size=14, color=MUTED)


def s13_detection(prs):
    s = base(prs, 13, "Detection", "It learns the recording, not the drone")
    text(s, 0.55, 1.42, 12.2, 0.3,
         "AUROC of the one-microphone detector. 0.5 is chance.",
         size=15, color=BODY)
    picture(s, DOCS / "analysis_detector.png", 0.45, 1.9, w=6.6)
    card(s, 7.35, 1.9, 5.4, 3.75)
    text(s, 7.58, 2.05, 5.0, 0.3, "WHAT ELSE WE TRIED", size=12, color=MUTED, bold=True)
    tried = [
        ("+ ESC-50", "0.49 → 0.59", "Only on the training-site flights."),
        ("+ DEMAND", "−0.03", "Worst held-out fold."),
        ("U-Net presence head", "0.73 / 0.59", "All 24 / device 3, same afternoon."),
    ]
    for i, (name, value, body) in enumerate(tried):
        y = 2.5 + i * 1.02
        text(s, 7.58, y, 3.0, 0.32, name, size=15, color=INK, bold=True)
        text(s, 10.3, y, 2.25, 0.32, value, size=15, color=INK, bold=True, align=PP_ALIGN.RIGHT)
        text(s, 7.58, y + 0.36, 5.0, 0.5, body, size=13, color=BODY)
    bottom_line(s, 5.95, "The detector studies the environment, not the drone. "
                "A new site needs quiet and drone recordings made there.")


def s14_close(prs):
    s = base(prs, 14, "Conclusion", "Time can replace microphones")
    found = [
        ("History stands in for microphones",
         "Nine frames let one recorder of six reach the full 24-microphone array."),
        ("The filter is an alternative",
         "It lifts single-frame models, but adds nothing on top of temporal context."),
        ("Detection is the open problem",
         "Drone vs no-drone does not generalize. Few negatives likely bias the numbers."),
    ]
    future = [
        ("More data", "Longer, more varied ambient noise and clean drone recordings."),
        ("A better array", "Length and geometry matched to the rotor wavelengths."),
        ("Do the gains hold?", "Check history and the filter as base accuracy rises."),
        ("A fair benchmark", "Balanced, multi-environment data for an unbiased score."),
    ]
    text(s, 0.55, 1.42, 6.0, 0.3, "WHAT WE FOUND", size=13, color=GREEN, bold=True)
    for i, (title, body) in enumerate(found):
        y = 1.8 + i * 1.6
        text(s, 0.55, y, 5.9, 0.34, title, size=18, color=INK, bold=True)
        text(s, 0.55, y + 0.42, 5.9, 0.9, body, size=15, color=BODY, line=1.06)
    card(s, 6.85, 1.42, 5.9, 5.4, fill=DARK, line=None)
    text(s, 7.1, 1.62, 5.4, 0.3, "FUTURE WORK", size=13, color="60A5FA", bold=True)
    for i, (title, body) in enumerate(future):
        y = 2.08 + i * 1.18
        text(s, 7.1, y, 5.4, 0.34, title, size=17, color="F5F5F4", bold=True)
        text(s, 7.1, y + 0.4, 5.4, 0.7, body, size=14, color="D6D3D1", line=1.06)


def overflow(prs) -> int:
    sw, sh = prs.slide_width, prs.slide_height
    n = 0
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            right = shape.left + shape.width
            bottom = shape.top + shape.height
            if right > sw + 20000 or bottom > sh + 20000:
                n += 1
                print(f"overflow slide {i}: right {right / 914400:.2f}in bottom {bottom / 914400:.2f}in")
    return n


def main() -> int:
    charts = save_charts()
    save_unet(DOCS / "unet_structure.png")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(SW * 914400)), Emu(int(SH * 914400))
    s01_title(prs)
    s02_problem(prs)
    s03_data(prs)
    s04_hardware(prs)
    s05_map(prs)
    s06_model(prs)
    s07_network(prs)
    s08_training(prs, charts)
    s09_paper(prs)
    s10_history(prs)
    s11_filter_gain(prs)
    s12_filter(prs)
    s13_detection(prs)
    s14_close(prs)
    if len(prs.slides) != TOTAL:
        raise SystemExit(f"{len(prs.slides)} slides, expected {TOTAL}")
    n = overflow(prs)
    if n:
        raise SystemExit(f"{n} shapes past the slide edge")
    prs.save(OUT)
    charts["training"].unlink()
    print(OUT)
    return 0


if __name__ == "__main__":
    sys.exit(main())
